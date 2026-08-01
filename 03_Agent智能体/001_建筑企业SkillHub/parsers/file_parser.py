# 解析器
from pathlib import Path



def parse_file(file_path):

    path = Path(file_path)

    suffix = path.suffix.lower()


    print("检测文件格式:", suffix)



    if suffix in [".doc", ".docx"]:

        return parse_word(file_path)



    elif suffix == ".pdf":

        return parse_pdf(file_path)



    elif suffix in [".xls", ".xlsx"]:

        return parse_excel(file_path)



    elif suffix == ".mpp":

        return parse_project(file_path)



    else:

        raise Exception(
            f"暂不支持文件格式:{suffix}"
        )





def parse_word(file_path):

    from docx import Document


    doc = Document(file_path)

    text=""


    for p in doc.paragraphs:

        text += p.text+"\n"


    return text






def parse_pdf(file_path):

    import fitz


    doc = fitz.open(file_path)

    text=""


    for page in doc:

        text += page.get_text()


    return text






def parse_excel(file_path):

    import openpyxl


    wb = openpyxl.load_workbook(
        file_path,
        data_only=True
    )


    text=""


    for sheet in wb:

        text += f"\n工作表:{sheet.title}\n"


        for row in sheet.iter_rows():

            values=[]


            for cell in row:

                if cell.value:

                    values.append(
                        str(cell.value)
                    )


            if values:

                text += " ".join(values)+"\n"



    return text






def parse_project(file_path):

    """
    Microsoft Project .mpp
    初版占位
    """

    return (
        "暂未解析MPP文件，"
        "后续接入Project解析模块"
    )